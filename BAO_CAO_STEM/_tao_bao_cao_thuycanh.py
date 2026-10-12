# -*- coding: utf-8 -*-
"""Bao cao STEM Thuy Canh IoT — nam hoc 2025-2026, ngay thang nam 2026."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

BASE = Path(r"C:\Users\Duong Phung\Downloads\thuycanh\BAO_CAO_STEM")
IMG = BASE / "hinh_minh_hoa"
OUT = BASE
OUT.mkdir(parents=True, exist_ok=True)


def set_run_font(run, size=13, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_para(doc, text, size=13, bold=False, italic=False, align="justify", space_after=6, space_before=0):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_bullets(doc, items, size=13, space_after=2):
    for item in items:
        add_para(doc, item, size=size, align="left", space_after=space_after)


def add_image(doc, path: Path, caption: str, width_cm=13.0):
    if not path.exists():
        add_para(doc, f"[Thiếu ảnh: {path.name} — có thể chèn thêm ảnh thật sau]", size=11, align="center")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    add_para(doc, caption, size=11, italic=True, align="center", space_after=10)


def set_cell_text(cell, text, bold=False, size=11):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)


def create_phu_luc_5():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2)

    # ===== COVER =====
    add_para(doc, "SỞ GIÁO DỤC VÀ ĐÀO TẠO HẢI PHÒNG", 14, True, align="center", space_after=2)
    add_para(doc, "TRƯỜNG THPT CÁT HẢI", 13, True, align="center", space_after=12)
    add_para(doc, "CUỘC THI KHOA HỌC KĨ THUẬT CẤP THÀNH PHỐ", 14, True, align="center", space_after=2)
    add_para(doc, "DÀNH CHO HỌC SINH TRUNG HỌC NĂM HỌC 2025 - 2026", 13, True, align="center", space_after=14)

    add_para(doc, "BÁO CÁO TÓM TẮT KẾT QUẢ NGHIÊN CỨU", 15, True, align="center", space_after=8)
    add_para(
        doc,
        "Tên dự án: Hệ thống thủy canh IoT STEM – Giám sát thông minh, trồng cây dễ dàng",
        13,
        True,
        align="left",
    )
    add_para(doc, "Loại dự án: Dự án khoa học (  )     Dự án kĩ thuật (☑)", 13, align="left")
    add_para(
        doc,
        "Lĩnh vực: Công nghệ thông tin – Internet vạn vật (IoT) / Nông nghiệp công nghệ cao / Giáo dục STEM liên ngành",
        13,
        align="left",
    )
    add_para(doc, "Mã dự án: ..........          Vị trí: ..........", 13, align="left", space_after=6)
    add_para(doc, "Học sinh thực hiện: ................................................", 13, align="left")
    add_para(doc, "Giáo viên hướng dẫn: Bùi Thị Diễm Hằng – môn Sinh học", 13, align="left", space_after=8)
    add_para(doc, "Hải Phòng, ngày ...... tháng ...... năm 2026", 13, align="center", space_after=14)

    add_para(doc, "MỤC LỤC", 13, True, align="center", space_after=6)
    for t in [
        "1. Vấn đề nghiên cứu",
        "2. Thiết kế và phương pháp",
        "    2.1. Cơ sở khoa học và yêu cầu thiết kế",
        "    2.2. Thiết kế tổng thể hệ thống",
        "    2.3. Thiết kế phần cứng và sơ đồ kết nối",
        "    2.4. Thiết kế phần mềm firmware và ứng dụng",
        "    2.5. Phương pháp nghiên cứu",
        "3. Thực hiện",
        "    3.1. Quy trình chế tạo phần cứng",
        "    3.2. Lập trình firmware ESP32",
        "    3.3. Ứng dụng di động theo dõi realtime",
        "    3.4. Kết quả thử nghiệm và hình ảnh minh họa",
        "4. Kết luận (tính mới, khả năng ứng dụng)",
        "5. Hướng phát triển: AI tối ưu nước và đèn tiết kiệm năng lượng",
        "6. Tài liệu tham khảo",
        "7. Phụ lục",
    ]:
        add_para(doc, t, 12, align="left", space_after=2)

    # ===== I =====
    add_para(doc, "I. VẤN ĐỀ NGHIÊN CỨU", 14, True, align="left", space_before=14)
    add_para(
        doc,
        "Trong những năm gần đây, nhu cầu rau sạch và xu hướng trồng cây tại nhà, tại trường ngày càng tăng. "
        "Thủy canh (hydroponics) – phương pháp trồng cây không dùng đất, cung cấp dinh dưỡng qua dung dịch – "
        "được xem là giải pháp phù hợp với không gian nhỏ như ban công, sân thượng hay phòng thực hành STEM. "
        "So với trồng đất truyền thống, thủy canh giúp kiểm soát tốt hơn môi trường rễ, tiết kiệm nước và giảm sâu bệnh từ đất. "
        "Tuy nhiên, để cây phát triển ổn định, người trồng phải theo dõi liên tục nhiều thông số quan trọng.",
    )
    add_para(
        doc,
        "Các thông số then chốt trong hệ thủy canh gồm: nhiệt độ và độ ẩm không khí (ảnh hưởng hô hấp và thoát hơi nước của cây); "
        "độ pH dung dịch dinh dưỡng (ảnh hưởng khả năng hấp thụ khoáng); nồng độ chất hòa tan TDS/EC (phản ánh mức dinh dưỡng); "
        "mực nước trong bồn (tránh cạn hoặc tràn); và cường độ ánh sáng (quyết định quang hợp). Nếu thiếu giám sát, "
        "chỉ sau thời gian ngắn cây có thể thiếu nước, lệch pH, thiếu sáng hoặc thừa nhiệt – dẫn đến vàng lá, chậm lớn, thậm chí chết cây.",
        space_before=4,
    )
    add_para(
        doc,
        "Trên thực tế, người mới bắt đầu và học sinh thường gặp các khó khăn sau:",
        space_before=4,
    )
    add_bullets(
        doc,
        [
            "- Phải kiểm tra thủ công nhiều lần trong ngày, mất thời gian và dễ quên.",
            "- Khó biết chính xác khi nào cần bổ sung nước, khi nào nên bật đèn trồng.",
            "- Hệ thống thủy canh thương mại hoàn chỉnh thường đắt, khó sửa, ít minh bạch về nguyên lý – không phù hợp làm mô hình học tập.",
            "- Mô hình DIY ở trường thường dừng ở “khung trồng + ống nước”, chưa gắn cảm biến, điều khiển tự động và theo dõi từ xa.",
            "- Thiếu cầu nối giữa kiến thức Sinh học (nhu cầu cây) với Tin học – Điện tử (IoT), khiến bài học STEM chưa đủ trực quan.",
        ],
    )
    add_para(
        doc,
        "Từ những vấn đề trên, nhóm đặt câu hỏi nghiên cứu: Có thể thiết kế và chế tạo một hệ thống thủy canh IoT "
        "giá thành hợp lý, dùng linh kiện phổ biến, vừa trồng cây được, vừa giám sát realtime qua điện thoại, "
        "vừa tự động hỗ trợ bơm tuần hoàn và đèn trồng – đồng thời trở thành mô hình học tập STEM liên ngành "
        "cho học sinh trung học hay không?",
        space_before=6,
    )
    add_para(
        doc,
        "Dự án “Hệ thống thủy canh IoT STEM – Giám sát thông minh, trồng cây dễ dàng” được thực hiện trong năm học "
        "2025 – 2026 nhằm trả lời câu hỏi trên bằng một sản phẩm kỹ thuật hoàn chỉnh: phần cứng ESP32 + cảm biến + "
        "bơm/đèn, firmware nhúng, ứng dụng Flutter Android/iOS, và định hướng phát triển AI tối ưu năng lượng.",
        space_before=4,
    )

    # ===== II =====
    add_para(doc, "II. THIẾT KẾ VÀ PHƯƠNG PHÁP NGHIÊN CỨU", 14, True, align="left", space_before=12)

    add_para(doc, "2.1. Cơ sở khoa học và yêu cầu thiết kế", 13, True, align="left")
    add_para(
        doc,
        "Về Sinh học – nông nghiệp: rau ăn lá trong hệ thủy canh cần môi trường rễ ẩm ổn định, dung dịch có pH phù hợp "
        "(thường khoảng trung tính hơi acid tùy loại cây), dinh dưỡng đủ nhưng không quá đặc, và ánh sáng đủ cho quang hợp. "
        "Về kỹ thuật: hệ IoT cần đo được các đại lượng vật lý/hóa học liên quan, truyền dữ liệu ổn định, điều khiển chấp hành "
        "an toàn, giao diện người dùng dễ hiểu. Về giáo dục STEM: sản phẩm phải “mở” – học sinh nhìn thấy cảm biến, "
        "hiểu được ngưỡng, giải thích được vì sao bơm/đèn bật tắt.",
    )
    add_para(doc, "Yêu cầu thiết kế cụ thể của dự án:", space_before=4)
    add_bullets(
        doc,
        [
            "- Đo được: nhiệt độ, độ ẩm, pH, TDS (tùy chọn), mực nước, ánh sáng.",
            "- Hiển thị tại chỗ trên màn hình thiết bị và từ xa trên điện thoại.",
            "- Tự động hóa cơ bản: bơm theo mực nước; đèn theo ánh sáng môi trường (AUTO).",
            "- Linh kiện dễ tìm, giá hợp lý, sơ đồ đấu nối rõ ràng.",
            "- App hỗ trợ đăng nhập, chọn máy theo mã chip, xem realtime và trạng thái hệ thống.",
            "- Có khả năng mở rộng sang AI tối ưu điện – nước trong giai đoạn phát triển.",
        ],
    )

    add_para(doc, "2.2. Thiết kế tổng thể hệ thống", 13, True, align="left", space_before=8)
    add_para(
        doc,
        "Hệ thống được tổ chức theo kiến trúc 3 lớp bổ trợ nhau:",
    )
    add_bullets(
        doc,
        [
            "Lớp 1 – Cảm biến và chấp hành: DHT22, LDR, pH, HC-SR04, TDS (tùy chọn), bơm 5V, đèn LED trồng, relay 2 kênh.",
            "Lớp 2 – Điều khiển nhúng: ESP32 (WiFi), firmware Arduino, màn hình TFT/OLED, WiFi portal cấu hình mạng, MQTT WSS.",
            "Lớp 3 – Ứng dụng người dùng: app Flutter “THỦY CANH” trên Android/iOS, đăng nhập backend AloT, theo dõi realtime theo chipId 790.",
        ],
    )
    add_para(
        doc,
        "Luồng dữ liệu: cảm biến → ESP32 xử lý/lọc nhiễu → hiển thị local + publish MQTT → app subscribe và vẽ dashboard. "
        "Luồng điều khiển: ESP32 tự quyết định AUTO theo ngưỡng; (hướng phát triển) app/AI gửi lịch tối ưu xuống thiết bị qua lệnh MQTT, "
        "nhưng vẫn giữ ràng buộc an toàn chống cạn nước.",
        space_before=4,
    )
    add_image(
        doc,
        IMG / "01_infographic_he_thong.jpg",
        "Hình 1. Infographic tổng quan hệ thống thủy canh IoT STEM: cảm biến, sơ đồ kết nối ESP32 và theo dõi qua điện thoại",
        width_cm=14.5,
    )

    add_para(doc, "2.3. Thiết kế phần cứng và sơ đồ kết nối", 13, True, align="left", space_before=8)
    add_para(
        doc,
        "Mô hình cơ khí dùng thùng nhựa trong suốt làm bồn dinh dưỡng, nắp trồng khoan nhiều lỗ đặt rọ net pot, "
        "khung ống PVC đỡ đèn LED và cố định cảm biến. Hộp điện bên hông chứa ESP32, relay và mạch nguồn. "
        "Cảm biến siêu âm HC-SR04 gắn hướng xuống mặt nước để đo mực nước không tiếp xúc – phù hợp môi trường ẩm và dễ bảo trì.",
    )
    add_image(
        doc,
        IMG / "06_prototype_den_led.jpg",
        "Hình 2. Nguyên mẫu hệ thống: khay trồng, khung PVC, đèn LED trồng và cảm biến gắn trên mô hình",
        width_cm=13,
    )
    add_image(
        doc,
        IMG / "07_prototype_bon_chua.jpg",
        "Hình 3. Bồn chứa dung dịch, nắp trồng nhiều lỗ, hộp điều khiển ESP32 và cảm biến siêu âm đo mực nước",
        width_cm=13,
    )

    add_para(doc, "Bảng phân công linh kiện và vai trò:", 13, True, align="left", space_before=6)
    add_bullets(
        doc,
        [
            "- DHT22: đo nhiệt độ và độ ẩm không khí quanh vùng cây; giao tiếp 1 dây dữ liệu, dễ dùng.",
            "- LDR (module digital DO): phát hiện đủ sáng / thiếu sáng để điều khiển đèn AUTO.",
            "- Cảm biến pH analog: theo dõi độ acid/kiềm dung dịch dinh dưỡng.",
            "- TDS (tùy chọn): ước lượng nồng độ chất hòa tan (ppm) – chỉ thị mức dinh dưỡng.",
            "- HC-SR04: đo khoảng cách tới mặt nước → suy ra mực nước % và trạng thái FULL/OK/LOW.",
            "- Bơm tuần hoàn 5V: đảo/trộn dung dịch, hỗ trợ oxy và đồng đều dinh dưỡng.",
            "- Đèn LED trồng: bổ sung ánh sáng khi môi trường tối.",
            "- Relay 2 kênh: cách ly điều khiển tải bơm và đèn khỏi chân GPIO.",
            "- ESP32: não bộ hệ thống – đọc cảm biến, ra quyết định, kết nối WiFi/MQTT, hiển thị.",
        ],
    )

    add_para(doc, "Bảng chân kết nối ESP32:", 13, True, align="left", space_before=6)
    table = doc.add_table(rows=9, cols=3)
    table.style = "Table Grid"
    for i, h in enumerate(["Thành phần", "Chân ESP32", "Ghi chú"]):
        set_cell_text(table.rows[0].cells[i], h, bold=True, size=11)
    for r, row in enumerate(
        [
            ("DHT22", "GPIO 4", "Nhiệt độ, độ ẩm không khí"),
            ("LDR (DO)", "GPIO 34", "Cảm quang digital 0/1"),
            ("pH analog", "GPIO 35", "Phân áp nếu tín hiệu 5V"),
            ("TDS (tùy chọn)", "GPIO 32", "ppm dinh dưỡng"),
            ("HC-SR04", "TRIG 12, ECHO 13", "ECHO cần phân áp 5V→3.3V"),
            ("Màn hình TFT SPI", "21, 22, 5, 16, 17, 15", "Hiển thị thông số tại chỗ"),
            ("Relay bơm", "GPIO 26", "Bơm tuần hoàn AUTO theo mực nước"),
            ("Relay đèn", "GPIO 27", "Đèn trồng AUTO theo ánh sáng"),
        ],
        start=1,
    ):
        for c, val in enumerate(row):
            set_cell_text(table.rows[r].cells[c], val, size=11)
    add_para(doc, "", space_after=4)

    add_para(doc, "Logic tự động hóa (rule-based) phiên bản hiện tại:", 13, True, align="left", space_before=4)
    add_bullets(
        doc,
        [
            "- Mực nước: khoảng cách cảm biến tới mặt nước ≤ 11 cm → FULL (đầy) → bơm tuần hoàn ON trong AUTO; ≥ 18 cm → LOW (thấp) → tắt bơm + cảnh báo cần đổ nước; giữa 11–18 cm → OK.",
            "- Đèn trồng AUTO theo LDR: môi trường tối → bật đèn; đủ sáng → tắt đèn để tiết kiệm điện.",
            "- App chủ yếu giám sát; máy tự chỉnh theo cảm biến khi ở chế độ AUTO.",
            "- Lọc nhiễu HC-SR04: lấy burst nhiều mẫu, median kết hợp Kalman 1D để ổn định số đo mực nước, giảm cảnh báo giả.",
        ],
    )

    add_para(doc, "2.4. Thiết kế phần mềm firmware và ứng dụng", 13, True, align="left", space_before=8)
    add_para(
        doc,
        "Firmware ThuyCanh_ESP32.ino chạy trên ESP32 Dev Module. Thiết bị mở WiFi portal AP tên “ThuyCanh” "
        "(truy cập http://192.168.4.1) để cấu hình mạng. Sau khi có Internet, ESP32 kết nối MQTT qua WebSocket SSL "
        "tới broker AloT, sử dụng chipId = 790 (không trùng máy khác). Dữ liệu publish theo topic tele/<device>/status "
        "với payload JSON dạng {\"value\": ...}.",
    )
    add_para(doc, "Các kênh dữ liệu chính:", space_before=4)
    add_bullets(
        doc,
        [
            "- 790: trạng thái online/offline của máy.",
            "- 790_temp / 790_humi: nhiệt độ (°C), độ ẩm (%).",
            "- 790_ph / 790_tds: pH và TDS (ppm).",
            "- 790_water / 790_dist / 790_water_alert: mực nước (%), khoảng cách (cm), cảnh báo FULL/LOW/OK.",
            "- 790_light: mức sáng (%).",
            "- 790_pump / 790_lamp: trạng thái bơm và đèn.",
            "- Lệnh điều khiển (mở rộng): cmnd/790_power|pump|lamp/POWER với ON/OFF/AUTO.",
        ],
    )
    add_para(
        doc,
        "Ứng dụng di động viết bằng Flutter, hỗ trợ Android và iOS. Người dùng đăng ký/đăng nhập qua API AloT, "
        "nhập mã chip gắn trên máy (ví dụ 790), rồi xem dashboard realtime. Giao diện dùng tông màu tối – xanh lá, "
        "phù hợp chủ đề nông nghiệp/IoT, nhấn mạnh các chỉ số quan trọng và trạng thái hệ thống.",
        space_before=4,
    )

    add_para(doc, "2.5. Phương pháp nghiên cứu", 13, True, align="left", space_before=8)
    add_bullets(
        doc,
        [
            "- Phương pháp nghiên cứu tài liệu: thủy canh cơ bản; vai trò pH, TDS, ánh sáng; ESP32; MQTT; lọc tín hiệu cảm biến.",
            "- Phương pháp thiết kế kỹ thuật: phân tích yêu cầu → chọn linh kiện → vẽ sơ đồ → lập trình → tích hợp.",
            "- Phương pháp thực nghiệm: lắp mô hình thật, hiệu chỉnh cảm biến, thử WiFi/MQTT, kiểm thử app trên điện thoại.",
            "- Phương pháp STEM liên ngành: Sinh học (nhu cầu cây) – Vật lý (đo lường) – Toán (ngưỡng, %, lọc nhiễu) – Tin học (nhúng + app).",
            "- Phương pháp đánh giá: đối chiếu số liệu màn hình thiết bị với app; quan sát phản ứng bơm/đèn khi thay đổi mực nước và ánh sáng; ghi nhận độ ổn định kết nối.",
        ],
    )

    # ===== III =====
    add_para(doc, "III. THỰC HIỆN", 14, True, align="left", space_before=12)

    add_para(doc, "3.1. Quy trình chế tạo phần cứng", 13, True, align="left")
    add_bullets(
        doc,
        [
            "Bước 1 – Cơ khí: chuẩn bị thùng chứa, khoan lỗ nắp trồng, dựng khung PVC và bố trí nguồn 5V an toàn.",
            "Bước 2 – Cảm biến: lắp DHT22, LDR gần vùng cây; gắn HC-SR04 hướng mặt nước; đặt đầu dò pH vào dung dịch; (nếu có) TDS.",
            "Bước 3 – Điện tử: đấu dây theo bảng GPIO; dùng GND chung; phân áp cho chân 5V (ECHO, có thể cả pH); nối relay bơm/đèn.",
            "Bước 4 – Hiển thị và chiếu sáng: lắp màn hình, đèn LED trên khung, đi dây gọn bằng dây rút.",
            "Bước 5 – Cấu hình: nạp firmware, kết nối WiFi qua portal AP “ThuyCanh”.",
            "Bước 6 – Hiệu chỉnh: kiểm tra từng cảm biến, đối chiếu mực nước thật với % hiển thị, xác nhận AUTO bơm/đèn đúng ngưỡng.",
        ],
    )

    add_para(doc, "3.2. Lập trình firmware ESP32", 13, True, align="left", space_before=8)
    add_para(
        doc,
        "Quá trình lập trình firmware tập trung vào độ ổn định và dễ bảo trì. Các khối chức năng gồm: đọc cảm biến định kỳ; "
        "lọc nhiễu siêu âm; cập nhật trạng thái FULL/OK/LOW; điều khiển relay theo AUTO; vẽ UI lên màn hình; duy trì MQTT "
        "và publish/subscribe topic. Khi mất mạng, thiết bị vẫn có thể duy trì logic an toàn local (đặc biệt với mực nước), "
        "và app sẽ hiển thị trạng thái offline để người dùng biết.",
    )

    add_para(doc, "3.3. Ứng dụng di động theo dõi realtime", 13, True, align="left", space_before=8)
    add_para(
        doc,
        "App “THỦY CANH” được thiết kế theo các màn hình chính: đăng nhập/đăng ký; chọn máy bằng mã chip; dashboard "
        "mực nước và môi trường trồng; khu vực trạng thái bơm & đèn tự động. Người dùng nhìn nhanh được hệ thống đang "
        "ổn hay cần can thiệp (ví dụ nước thấp, máy offline).",
    )
    add_para(doc, "Minh họa giao diện ứng dụng:", 13, True, align="left", space_before=6)
    add_image(doc, IMG / "03_app_dang_nhap.png", "Hình 4. Màn hình đăng nhập app THỦY CANH", width_cm=7.5)
    add_image(doc, IMG / "02_app_chon_may.png", "Hình 5. Chọn máy theo tên chip ESP32 (máy 790)", width_cm=7.5)
    add_image(
        doc,
        IMG / "04_app_dashboard.png",
        "Hình 6. Dashboard realtime: mực nước, nhiệt độ, độ ẩm, pH, TDS",
        width_cm=7.5,
    )
    add_image(
        doc,
        IMG / "05_app_bom_den.png",
        "Hình 7. Trạng thái tự động hóa bơm theo mực nước và đèn theo ánh sáng",
        width_cm=7.5,
    )

    add_para(doc, "3.4. Kết quả thử nghiệm", 13, True, align="left", space_before=8)
    add_para(
        doc,
        "Sau quá trình lắp đặt, lập trình và chạy thử trong năm học 2025 – 2026, hệ thống đạt các kết quả chính sau:",
    )
    add_bullets(
        doc,
        [
            "- Đọc và hiển thị được các thông số: nhiệt độ, độ ẩm, pH, TDS, mực nước, ánh sáng trên thiết bị và trên app.",
            "- Kết nối WiFi/MQTT hoạt động; app phản ánh trạng thái online/offline của máy 790.",
            "- Bơm tự chỉnh theo ngưỡng mực nước; đèn tự tắt khi đủ sáng, bật khi thiếu sáng (AUTO).",
            "- Mô hình vật lý hoàn chỉnh: bồn dinh dưỡng, khay trồng nhiều lỗ, khung đèn LED, hộp điều khiển gọn.",
            "- Học sinh quan sát được chuỗi nhân quả: thay đổi môi trường → cảm biến → thuật toán ngưỡng → hành động bơm/đèn → số liệu trên điện thoại.",
            "- Tạo được bộ tài liệu kỹ thuật (sơ đồ chân, topic MQTT, ảnh hệ thống/app) phục vụ thuyết trình cuộc thi.",
        ],
    )
    add_para(
        doc,
        "Hạn chế còn tồn tại (để định hướng cải tiến): độ chính xác cảm biến pH/TDS phụ thuộc hiệu chỉnh; "
        "kết nối phụ thuộc WiFi; chế độ AUTO hiện dùng ngưỡng cố định chưa tự học theo mùa/thời tiết. "
        "Các hạn chế này được xử lý một phần trong hướng phát triển AI ở Mục V.",
        space_before=4,
    )
    add_para(
        doc,
        "Ghi chú: Có thể chèn thêm ảnh thật quá trình lắp ráp, ảnh cây đang trồng, ảnh màn hình thiết bị và video demo vào Phụ lục khi nộp hồ sơ.",
        italic=True,
        space_before=4,
    )

    # ===== IV =====
    add_para(doc, "IV. KẾT LUẬN", 14, True, align="left", space_before=12)
    add_para(doc, "1. Tính mới và sáng tạo", 13, True, align="left")
    add_bullets(
        doc,
        [
            "- Kết hợp mô hình thủy canh thực tế với hệ IoT hoàn chỉnh (cảm biến – ESP32 – MQTT – app) trong một sản phẩm STEM học đường.",
            "- Dùng linh kiện phổ biến, giá hợp lý; sơ đồ rõ; phù hợp học sinh trung học tự làm và bảo trì.",
            "- Có đồng thời hiển thị tại chỗ và giám sát từ xa; có tự động hóa bơm/đèn chứ không chỉ đo lường thụ động.",
            "- Ứng dụng lọc tín hiệu (median + Kalman) cho siêu âm nhằm tăng độ ổn định mực nước.",
            "- App đa nền tảng (Android/iOS), định danh thiết bị bằng chipId, có thể mở rộng nhiều máy.",
            "- Định hướng nâng cấp AI tối ưu điện – nước tạo chiều sâu phát triển vượt mô hình IoT ngưỡng cứng thông thường.",
        ],
    )

    add_para(doc, "2. Khả năng ứng dụng và phát triển", 13, True, align="left", space_before=6)
    add_bullets(
        doc,
        [
            "- Ứng dụng ngay làm mô hình học STEM tại câu lạc bộ, phòng thực hành Tin – Sinh – Công nghệ.",
            "- Trồng rau ăn lá quy mô nhỏ tại nhà/ban công với khả năng theo dõi bằng điện thoại.",
            "- Nhân rộng: thêm máy mới chỉ cần đổi chipId và cấu hình WiFi.",
            "- Hướng phát triển ngắn hạn: thông báo đẩy; biểu đồ lịch sử; camera; nhiều tầng trồng.",
            "- Hướng phát triển trọng tâm năm 2026 trở đi: module AI tối ưu nước và đèn tiết kiệm năng lượng (Mục V).",
        ],
    )

    # ===== V AI =====
    add_para(doc, "V. HƯỚNG PHÁT TRIỂN: ỨNG DỤNG AI TỐI ƯU NƯỚC VÀ ĐÈN TIẾT KIỆM NĂNG LƯỢNG", 14, True, align="left", space_before=12)
    add_para(
        doc,
        "Phiên bản hiện tại điều khiển bơm và đèn theo ngưỡng cố định (rule-based). Cách này đơn giản, an toàn và phù hợp giai đoạn học IoT cơ bản. "
        "Tuy nhiên, ngưỡng cố định chưa tận dụng hết dữ liệu lịch sử để thích nghi theo thời tiết, mùa vụ và giai đoạn sinh trưởng. "
        "Hướng phát triển tiếp theo trong và sau năm học 2025 – 2026 là bổ sung lớp Trí tuệ nhân tạo (AI) vào ứng dụng/đám mây "
        "để tối ưu hai tác nhân tiêu tốn năng lượng chính: bơm tuần hoàn và đèn trồng LED.",
    )

    add_para(doc, "5.1. Vì sao cần AI tối ưu năng lượng?", 13, True, align="left", space_before=6)
    add_bullets(
        doc,
        [
            "- Bơm chạy dày hoặc kéo dài gây tốn điện, rung ồn và có thể làm nóng dung dịch không cần thiết.",
            "- Đèn LED bật khi ban ngày đã đủ nắng là lãng phí rõ rệt.",
            "- Nhu cầu cây thay đổi theo giai đoạn (gieo mầm, sinh trưởng, thu hoạch) và theo mùa.",
            "- AI có thể học từ dữ liệu cảm biến để đề xuất lịch vận hành “đủ nhu cầu cây – ít tốn điện/nước nhất”.",
        ],
    )

    add_para(doc, "5.2. Kiến trúc module AI đề xuất", 13, True, align="left", space_before=6)
    add_bullets(
        doc,
        [
            "1) Thu thập: nhiệt độ, độ ẩm, ánh sáng, mực nước, thời gian bơm/đèn ON, (tùy chọn) ước lượng điện năng.",
            "2) Lưu lịch sử theo giờ/ngày trên cloud hoặc trong app.",
            "3) Mô hình AI/thuật toán: dự báo nhu cầu tuần hoàn và lịch chiếu sáng tối ưu.",
            "4) Sinh lịch điều khiển thông minh gửi ESP32 qua MQTT; vẫn giữ ngưỡng an toàn LOW/FULL.",
            "5) App hiển thị gợi ý và thống kê tiết kiệm năng lượng bằng tiếng Việt dễ hiểu.",
        ],
    )

    add_para(doc, "5.3. AI tối ưu nước (bơm tuần hoàn)", 13, True, align="left", space_before=6)
    add_bullets(
        doc,
        [
            "- Mục tiêu: đủ tuần hoàn/oxy/dinh dưỡng cho rễ, giảm số lần chạy bơm thừa.",
            "- Đầu vào: chuỗi mực nước, nhiệt độ, chu kỳ bơm cũ, cảnh báo LOW/FULL.",
            "- Đầu ra: lịch bơm theo khung giờ, thời lượng mỗi lần, chế độ tiết kiệm / sinh trưởng nhanh.",
            "- Ràng buộc cứng: không để bơm chạy khô khi LOW; có thể giảm tần suất khi FULL.",
            "- Lộ trình thuật toán: thống kê + quy tắc thông minh → dự báo chuỗi thời gian → (nâng cao) học tăng cường khi đủ dữ liệu.",
        ],
    )

    add_para(doc, "5.4. AI tối ưu đèn chiếu sáng", 13, True, align="left", space_before=6)
    add_bullets(
        doc,
        [
            "- Mục tiêu: đủ ánh sáng cho quang hợp, cực tiểu hóa giờ đèn ON không cần thiết.",
            "- Đầu vào: LDR/ánh sáng, giờ trong ngày, nhiệt độ, (sau này) ảnh cây/giai đoạn sinh trưởng, dự báo thời tiết.",
            "- Đầu ra: lịch bật/tắt; khung giờ bổ sung LED; chế độ tận dụng nắng ban ngày.",
            "- Tiêu chí: đạt ngưỡng ánh sáng tích lũy tối thiểu cho loại rau, đồng thời giảm điện năng.",
        ],
    )

    add_para(doc, "5.5. Tích hợp trên ứng dụng và lợi ích kỳ vọng", 13, True, align="left", space_before=6)
    add_bullets(
        doc,
        [
            "- Tab “AI tiết kiệm năng lượng”: lịch bơm/đèn 24 giờ tới; nút Áp dụng / Chỉ xem gợi ý.",
            "- Thống kê % điện tiết kiệm theo tuần so với chế độ ngưỡng cứng.",
            "- Giải thích ngắn cho học sinh: vì sao AI đề xuất tắt đèn buổi trưa nắng, giảm bơm ban đêm…",
            "- Lợi ích: tiết kiệm điện – nước, tăng tuổi thọ thiết bị, nâng tầm STEM hướng nông nghiệp xanh thông minh.",
        ],
    )
    add_image(
        doc,
        IMG / "05_app_bom_den.png",
        "Hình 8. Nền tảng bơm/đèn AUTO hiện tại – sẽ nâng cấp bằng lớp AI tối ưu năng lượng",
        width_cm=7.5,
    )

    # ===== VI refs =====
    add_para(doc, "VI. TÀI LIỆU THAM KHẢO", 14, True, align="left", space_before=12)
    add_bullets(
        doc,
        [
            "1. Espressif Systems – ESP32 Technical Reference / Arduino core for ESP32.",
            "2. Adafruit – DHT sensor library; tài liệu HC-SR04 ultrasonic ranging.",
            "3. MQTT.org – MQTT protocol specification (tham khảo).",
            "4. Flutter.dev – Documentation.",
            "5. Tài liệu thủy canh cơ bản về pH, EC/TDS và ánh sáng cho rau ăn lá.",
            "6. Mã nguồn dự án: firmware/ThuyCanh_ESP32.ino và ứng dụng Flutter thư mục thuycanh/ (năm học 2025 – 2026).",
            "7. Các tài liệu mở về tối ưu năng lượng và ứng dụng AI trong nông nghiệp thông minh (định hướng phát triển).",
        ],
    )

    add_para(doc, "VII. PHỤ LỤC", 14, True, align="left", space_before=12)
    add_bullets(
        doc,
        [
            "- Phụ lục ảnh: Hình 1–8 (infographic, nguyên mẫu, app).",
            "- Phụ lục kỹ thuật: bảng GPIO, topic MQTT, ngưỡng mực nước 11 cm / 18 cm.",
            "- Phụ lục định hướng AI: sơ đồ dữ liệu → mô hình → lệnh MQTT (bổ sung khi triển khai).",
            "- (Bổ sung) Ảnh thật lắp ráp, ảnh cây, biên bản thử nghiệm, video demo.",
        ],
    )

    add_para(doc, "Hải Phòng, ngày ...... tháng ...... năm 2026", 13, align="right", space_before=16)
    add_para(doc, "Học sinh thực hiện ký tên", 13, True, align="right")

    out = OUT / "Phu_luc_5_Bao_cao_tom_tat_ThuyCanh_IoT.docx"
    doc.save(out)
    print("Saved", out)
    return out


def create_phu_luc_7():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.5)
        s.right_margin = Cm(2)

    add_para(doc, "TRƯỜNG THPT CÁT HẢI", 14, True, align="center", space_after=10)
    add_para(doc, "CUỘC THI KHOA HỌC KĨ THUẬT CẤP THÀNH PHỐ", 14, True, align="center", space_after=2)
    add_para(doc, "DÀNH CHO HỌC SINH TRUNG HỌC NĂM HỌC 2025 - 2026", 13, True, align="center", space_after=10)
    add_para(doc, "KẾ HOẠCH NGHIÊN CỨU", 16, True, align="center", space_after=12)

    add_para(
        doc,
        "Tên đề tài: Hệ thống thủy canh IoT STEM – Giám sát thông minh, trồng cây dễ dàng",
        13,
        True,
        align="left",
    )
    add_para(
        doc,
        "Lĩnh vực dự thi: Công nghệ thông tin / Khoa học ứng dụng – Nông nghiệp công nghệ / Giáo dục STEM",
        13,
        align="left",
    )
    add_para(doc, "Họ tên học sinh 1: ................................................", 13, align="left")
    add_para(doc, "Họ tên học sinh 2: ................................................", 13, align="left")
    add_para(doc, "Giáo viên hướng dẫn: Bùi Thị Diễm Hằng – Giảng dạy môn Sinh học", 13, align="left", space_after=8)
    add_para(doc, "Hải Phòng, tháng ...... năm 2026", 13, align="center", space_after=12)

    add_image(
        doc,
        IMG / "01_infographic_he_thong.jpg",
        "Hình minh họa định hướng đề tài: hệ thống thủy canh IoT STEM",
        width_cm=13.5,
    )

    # I
    add_para(doc, "I. GIỚI THIỆU VÀ TÍNH CẤP THIẾT CỦA ĐỀ TÀI", 14, True, align="left", space_before=8)
    add_para(doc, "1. Lý do chọn đề tài", 13, True, align="left")
    add_para(
        doc,
        "Năm học 2025 – 2026, phong trào giáo dục STEM và chuyển đổi số trong nhà trường tiếp tục được đẩy mạnh. "
        "Đồng thời, học sinh và gia đình quan tâm hơn đến rau sạch, kỹ năng tự làm và ứng dụng công nghệ vào đời sống. "
        "Thủy canh là mô hình trực quan để hiểu mối quan hệ giữa môi trường sống của cây và các đại lượng đo được. "
        "Khi gắn thêm IoT, mô hình không còn là “đồ chơi trồng cây” mà trở thành hệ giám sát – tự động hóa thật, "
        "gần với nông nghiệp công nghệ cao. Nhóm chọn đề tài để vừa giải quyết bài toán theo dõi dinh dưỡng/mực nước "
        "khi trồng thủy canh, vừa tạo sản phẩm học tập liên môn Sinh – Lý – Toán – Tin.",
    )
    add_para(
        doc,
        "Nếu chỉ dừng ở theo dõi thủ công, học sinh khó duy trì thói quen đo đạc và dễ bỏ cuộc khi cây gặp sự cố. "
        "Một hệ thống có cảm biến, app điện thoại và tự động hóa bơm/đèn sẽ giảm rào cản này, đồng thời mở ra câu hỏi "
        "nghiên cứu sâu hơn về tối ưu năng lượng bằng AI trong giai đoạn phát triển.",
        space_before=4,
    )

    add_para(doc, "2. Tính mới, tính cấp thiết", 13, True, align="left", space_before=6)
    add_bullets(
        doc,
        [
            "- Cấp thiết: thiếu giải pháp thủy canh IoT giá vừa túi tiền, dễ làm, phù hợp môi trường học đường năm học 2025 – 2026.",
            "- Tính mới: tích hợp đồng thời cảm biến môi trường + dinh dưỡng + mực nước + điều khiển bơm/đèn + app realtime.",
            "- Tính liên ngành rõ: Sinh học cây trồng – Điện tử nhúng – Lập trình app – Toán ngưỡng/lọc tín hiệu.",
            "- Chiều sâu phát triển: định hướng AI tối ưu nước và đèn tiết kiệm năng lượng, hướng tới nông nghiệp đô thị xanh.",
            "- Giá trị thực tiễn: dùng được ngay để trồng quy mô nhỏ và thuyết trình/demo tại trường.",
        ],
    )

    # II
    add_para(doc, "II. MỤC TIÊU VÀ GIẢ THUYẾT NGHIÊN CỨU", 14, True, align="left", space_before=10)
    add_para(doc, "1. Mục tiêu nghiên cứu", 13, True, align="left")
    add_para(doc, "Mục tiêu tổng quát:", bold=True, align="left", space_after=2)
    add_para(
        doc,
        "Thiết kế, chế tạo, lập trình và thử nghiệm hệ thống thủy canh IoT STEM có khả năng giám sát realtime và tự động "
        "hỗ trợ duy trì môi trường trồng (bơm, đèn), phục vụ học tập và ứng dụng quy mô nhỏ trong năm học 2025 – 2026.",
    )
    add_para(doc, "Mục tiêu cụ thể:", bold=True, align="left", space_before=4, space_after=2)
    add_bullets(
        doc,
        [
            "1) Lắp mô hình thủy canh quy mô nhỏ với ESP32 và cảm biến DHT22, LDR, pH, HC-SR04 (TDS tùy chọn).",
            "2) Viết firmware đọc cảm biến, hiển thị tại chỗ, điều khiển bơm/đèn AUTO theo ngưỡng an toàn.",
            "3) Xây dựng app Flutter (Android/iOS): đăng nhập, chọn chip, dashboard realtime, xem trạng thái bơm/đèn.",
            "4) Đánh giá độ ổn định kết nối, độ phản hồi tự động hóa và giá trị giáo dục STEM của mô hình.",
            "5) Xây dựng hồ sơ báo cáo, hình ảnh minh họa và định hướng phát triển AI tối ưu điện – nước.",
            "6) (Giai đoạn phát triển) Nghiên cứu tích hợp AI đề xuất lịch bơm và đèn nhằm tiết kiệm năng lượng.",
        ],
    )

    add_para(doc, "2. Giả thuyết nghiên cứu", 13, True, align="left", space_before=6)
    add_para(
        doc,
        "Giả thuyết 1: Nếu xây dựng hệ thống thủy canh gắn ESP32 với bộ cảm biến phù hợp, kết nối MQTT và ứng dụng di động, "
        "đồng thời thiết lập ngưỡng tự động cho bơm và đèn, thì có thể giám sát được các thông số quan trọng của môi trường trồng "
        "theo thời gian thực và giảm thao tác theo dõi thủ công.",
    )
    add_para(
        doc,
        "Giả thuyết 2: Nếu tổ chức sản phẩm theo hướng mở (nhìn thấy cảm biến, hiểu được ngưỡng, có app minh họa), "
        "thì mô hình sẽ hỗ trợ học tập STEM hiệu quả hơn so với mô hình trồng thủy canh không có IoT.",
        space_before=4,
    )
    add_para(
        doc,
        "Giả thuyết 3 (định hướng phát triển): Nếu bổ sung lớp AI học từ dữ liệu cảm biến và lịch sử vận hành, "
        "hệ thống có thể đề xuất lịch bơm – chiếu sáng thông minh hơn ngưỡng cố định, góp phần tối ưu nước và tiết kiệm điện năng.",
        space_before=4,
    )

    # III
    add_para(doc, "III. PHẠM VI NGHIÊN CỨU", 14, True, align="left", space_before=10)
    add_para(doc, "1. Địa điểm nghiên cứu", 13, True, align="left")
    add_para(
        doc,
        "Trường THPT Cát Hải (Hải Phòng) và/hoặc không gian thực hành tại nhà của học sinh thực hiện đề tài. "
        "Việc thử nghiệm app được thực hiện trên điện thoại Android/iOS của nhóm.",
    )
    add_para(doc, "2. Thời gian nghiên cứu", 13, True, align="left", space_before=4)
    add_para(
        doc,
        "Thực hiện trong năm học 2025 – 2026. Khung thời gian dự kiến các mốc chính trong năm 2026 "
        "(có thể điều chỉnh theo lịch cuộc thi và điều kiện nhà trường): nghiên cứu tài liệu – thiết kế – chế tạo – "
        "lập trình – thử nghiệm – hoàn thiện báo cáo – định hướng AI.",
    )
    add_para(doc, "3. Đối tượng và giới hạn nghiên cứu", 13, True, align="left", space_before=4)
    add_bullets(
        doc,
        [
            "- Đối tượng: mô hình thủy canh IoT quy mô nhỏ (desktop), trồng rau ăn lá; hệ cảm biến – ESP32 – app giám sát.",
            "- Giới hạn: tự động hóa cơ bản (bơm, đèn theo ngưỡng); chưa đi sâu điều khiển bơm dinh dưỡng đa kênh công nghiệp.",
            "- App ưu tiên giám sát realtime; điều khiển nâng cao và AI tối ưu là hướng phát triển tiếp theo.",
            "- Độ chính xác tuyệt đối kiểu phòng thí nghiệm không phải mục tiêu chính; ưu tiên tính ổn định, trực quan và giá trị học tập.",
        ],
    )

    # IV
    add_para(doc, "IV. PHƯƠNG PHÁP NGHIÊN CỨU", 14, True, align="left", space_before=10)
    add_para(
        doc,
        "Nhóm kết hợp nhiều phương pháp để vừa đảm bảo tính khoa học vừa phù hợp điều kiện học sinh trung học:",
    )
    add_bullets(
        doc,
        [
            "- Nghiên cứu lý thuyết: thủy canh; thông số pH/TDS/ánh sáng; kiến trúc IoT ESP32 – MQTT – app; cơ sở tối ưu năng lượng.",
            "- Thiết kế kỹ thuật: chọn linh kiện dễ tìm; lập sơ đồ đấu nối; thiết kế cơ khí mô hình; phân rã chức năng phần mềm.",
            "- Thực nghiệm chế tạo và lập trình: firmware Arduino; app Flutter; kiểm thử từng module rồi tích hợp toàn hệ.",
            "- Quan sát – đo đạc – ghi nhận: so sánh số liệu cảm biến với quan sát thực tế; ghi log trạng thái online/offline, bơm/đèn.",
            "- Phân tích – tổng hợp: đánh giá ưu nhược điểm, đề xuất cải tiến và xây dựng hướng phát triển AI.",
            "- Phương pháp STEM: mỗi quyết định kỹ thuật gắn với kiến thức môn học (ví dụ ngưỡng mực nước gắn Toán/Vật lý đo lường; pH gắn Sinh/Hóa).",
        ],
    )
    add_image(
        doc,
        IMG / "07_prototype_bon_chua.jpg",
        "Hình. Mô hình thực nghiệm thủy canh IoT đang chế tạo/thử nghiệm",
        width_cm=12,
    )
    add_image(
        doc,
        IMG / "06_prototype_den_led.jpg",
        "Hình. Mô hình với khung đèn LED trồng và cảm biến trên hệ thống",
        width_cm=12,
    )

    # V
    add_para(doc, "V. TIẾN ĐỘ THỰC HIỆN (NĂM HỌC 2025 – 2026 / CÁC MỐC NĂM 2026)", 14, True, align="left", space_before=10)
    add_para(
        doc,
        "Bảng tiến độ dưới đây dùng mốc theo giai/tháng trong năm 2026 để khớp thời điểm hoàn thiện hồ sơ. "
        "Các tuần đầu có thể bắt đầu từ cuối năm 2025 tùy lịch nhà trường.",
        italic=True,
        space_after=6,
    )
    prog = doc.add_table(rows=10, cols=3)
    prog.style = "Table Grid"
    for i, h in enumerate(["Thời gian", "Công việc", "Ghi chú"]):
        set_cell_text(prog.rows[0].cells[i], h, bold=True, size=11)
    schedule = [
        ("Giai đoạn 1 (đầu 2026)", "Tìm hiểu thủy canh, IoT ESP32, MQTT; chốt yêu cầu đề tài", "Lý thuyết"),
        ("Giai đoạn 2", "Thiết kế sơ đồ phần cứng, mua linh kiện, dựng khung mô hình", "Thiết kế – mua sắm"),
        ("Giai đoạn 3", "Lắp cảm biến, relay, màn hình; viết và nạp firmware", "Chế tạo – lập trình"),
        ("Giai đoạn 4", "Xây dựng app Flutter: đăng nhập, chọn chip, dashboard realtime", "Phần mềm"),
        ("Giai đoạn 5", "Tích hợp MQTT, hiệu chỉnh ngưỡng bơm/đèn, thử nghiệm ổn định", "Thử nghiệm"),
        ("Giai đoạn 6", "Trồng thử cây, ghi nhận kết quả, chụp ảnh/video minh họa", "Thực nghiệm Sinh"),
        ("Giai đoạn 7 (2026)", "Hoàn thiện báo cáo, kế hoạch, poster/infographic, luyện thuyết trình", "Hồ sơ cuộc thi"),
        ("Giai đoạn 8 (phát triển)", "Thiết kế module AI tối ưu bơm nước & đèn tiết kiệm năng lượng", "Hướng phát triển"),
        ("Thường xuyên", "Họp nhóm, báo cáo giáo viên hướng dẫn, cập nhật nhật ký nghiên cứu", "Quản lý tiến độ"),
    ]
    for r, row in enumerate(schedule, start=1):
        for c, val in enumerate(row):
            set_cell_text(prog.rows[r].cells[c], val, size=11)
    add_para(doc, "", space_after=6)

    # VI
    add_para(doc, "VI. DỰ KIẾN KẾT QUẢ VÀ KẾT LUẬN", 14, True, align="left", space_before=8)
    add_para(doc, "1. Kết quả dự kiến / đã đạt", 13, True, align="left")
    add_bullets(
        doc,
        [
            "- Mô hình thủy canh IoT hoạt động: đo nhiệt độ, độ ẩm, pH, TDS, mực nước, ánh sáng.",
            "- Tự động hóa bơm theo mực nước và đèn theo ánh sáng (AUTO).",
            "- Ứng dụng điện thoại theo dõi realtime, nhận diện máy theo chipId (790).",
            "- Bộ tài liệu kỹ thuật và ảnh minh họa hệ thống/app phục vụ thuyết trình năm học 2025 – 2026.",
            "- Định hướng rõ giai đoạn nâng cấp AI tối ưu nước và đèn nhằm tiết kiệm năng lượng.",
        ],
    )
    add_image(
        doc,
        IMG / "03_app_dang_nhap.png",
        "Hình. App đăng nhập – cửa ngõ sử dụng hệ thống giám sát",
        width_cm=7.0,
    )
    add_image(
        doc,
        IMG / "02_app_chon_may.png",
        "Hình. Chọn máy theo mã chip ESP32",
        width_cm=7.0,
    )
    add_image(
        doc,
        IMG / "04_app_dashboard.png",
        "Hình. Dashboard giám sát realtime trên điện thoại",
        width_cm=7.0,
    )
    add_image(
        doc,
        IMG / "05_app_bom_den.png",
        "Hình. Trạng thái bơm/đèn AUTO – nền tảng để phát triển AI tối ưu năng lượng",
        width_cm=7.0,
    )

    add_para(doc, "1.1. Hướng phát triển: ứng dụng AI tối ưu nước và đèn tiết kiệm năng lượng", 13, True, align="left", space_before=8)
    add_para(
        doc,
        "Sau khi hệ IoT vận hành ổn định, nhóm định hướng phát triển lớp AI trong ứng dụng để không chỉ “đo và bật/tắt theo ngưỡng”, "
        "mà còn học từ dữ liệu thực tế, tối ưu hai nguồn tiêu thụ chính: nước/tuần hoàn (bơm) và điện chiếu sáng (đèn LED).",
    )
    add_bullets(
        doc,
        [
            "- AI tối ưu nước: đề xuất chu kỳ bơm theo mực nước, nhiệt độ và giai đoạn cây; giảm chạy bơm thừa; giữ an toàn chống cạn.",
            "- AI tối ưu đèn: lập lịch bật/tắt theo ánh sáng môi trường và khung giờ hiệu quả; tận dụng nắng; bổ sung LED khi thiếu sáng.",
            "- Trên app: tab “AI tiết kiệm năng lượng”, áp dụng lịch AI, thống kê % điện tiết kiệm, giải thích gợi ý bằng tiếng Việt.",
            "- Lợi ích: tiết kiệm điện – nước, tăng tuổi thọ thiết bị, nâng tầm đề tài STEM theo hướng nông nghiệp thông minh bền vững.",
        ],
    )

    add_para(doc, "2. Kết luận", 13, True, align="left", space_before=6)
    add_para(
        doc,
        "Đề tài khả thi về kỹ thuật và phù hợp mục tiêu cuộc thi khoa học kỹ thuật dành cho học sinh trung học năm học 2025 – 2026. "
        "Sản phẩm vừa mang giá trị ứng dụng (trồng và giám sát thủy canh), vừa mang giá trị giáo dục STEM rõ rệt. "
        "Trong năm 2026, nhóm tiếp tục hoàn thiện độ chính xác cảm biến, ổn định kết nối, bổ sung ảnh/video thực nghiệm "
        "và triển khai dần module AI tối ưu bơm nước – đèn chiếu sáng để hệ thống vận hành tiết kiệm năng lượng, "
        "thông minh hơn và gần với mô hình nông nghiệp đô thị xanh.",
    )

    # VII
    add_para(doc, "VII. KINH PHÍ VÀ NGUỒN LỰC", 14, True, align="left", space_before=10)
    add_para(doc, "1. Kinh phí (ước tính, điều chỉnh theo giá thị trường năm 2026)", 13, True, align="left")
    cost = doc.add_table(rows=9, cols=3)
    cost.style = "Table Grid"
    for i, h in enumerate(["Hạng mục", "Đơn giá ước tính", "Ghi chú"]):
        set_cell_text(cost.rows[0].cells[i], h, bold=True, size=11)
    for r, row in enumerate(
        [
            ("ESP32 Dev Module", "80.000 – 150.000đ", "Vi điều khiển WiFi"),
            ("DHT22", "30.000 – 50.000đ", "Nhiệt độ, độ ẩm"),
            ("LDR module", "5.000 – 15.000đ", "Ánh sáng"),
            ("Cảm biến pH analog", "100.000 – 200.000đ", "Dinh dưỡng"),
            ("HC-SR04", "15.000 – 30.000đ", "Mực nước"),
            ("Bơm 5V + relay + đèn LED + cơ khí", "150.000 – 350.000đ", "Chấp hành + mô hình"),
            ("Chi phí phát triển AI/cloud (sau)", "tùy chọn", "API/hosting nếu dùng cloud AI"),
            ("Tổng ước tính giai đoạn IoT", "khoảng 500.000 – 900.000đ", "Tùy cấu hình màn hình/TDS"),
        ],
        start=1,
    ):
        for c, val in enumerate(row):
            set_cell_text(cost.rows[r].cells[c], val, size=11, bold=(r == 8))
    add_para(doc, "", space_after=6)

    add_para(doc, "2. Nguồn lực", 13, True, align="left")
    add_bullets(
        doc,
        [
            "- Nhân lực: nhóm học sinh thực hiện + giáo viên hướng dẫn (Sinh học).",
            "- Thiết bị: máy tính nạp code, điện thoại chạy app, bộ linh kiện IoT và mô hình thủy canh.",
            "- Phần mềm: Arduino IDE, Flutter SDK, broker MQTT AloT, GitHub (mã nguồn/build).",
            "- Hỗ trợ chuyên môn: giáo viên hướng dẫn; tài liệu mở về ESP32, MQTT, thủy canh, AI ứng dụng.",
        ],
    )

    add_para(doc, "VIII. TÀI LIỆU THAM KHẢO", 14, True, align="left", space_before=10)
    add_bullets(
        doc,
        [
            "1. Espressif Systems – ESP32 Technical Reference / Arduino core for ESP32.",
            "2. Adafruit – DHT sensor library; hướng dẫn HC-SR04 ultrasonic ranging.",
            "3. MQTT.org – MQTT Version specification (tham khảo giao thức).",
            "4. Flutter.dev – Documentation.",
            "5. Tài liệu thủy canh cơ bản về pH, EC/TDS và ánh sáng cho rau ăn lá.",
            "6. Mã nguồn và README dự án thuycanh (firmware + app) – năm học 2025 – 2026.",
            "7. Tài liệu tham khảo về nông nghiệp thông minh và tối ưu năng lượng bằng dữ liệu/AI (định hướng phát triển).",
        ],
    )

    add_para(doc, "Hải Phòng, ngày ...... tháng ...... năm 2026", 13, align="right", space_before=16)

    sign = doc.add_table(rows=2, cols=2)
    set_cell_text(sign.rows[0].cells[0], "GIÁO VIÊN HƯỚNG DẪN", bold=True, size=12)
    set_cell_text(sign.rows[0].cells[1], "HỌC SINH THỰC HIỆN", bold=True, size=12)
    set_cell_text(sign.rows[1].cells[0], "\n\n\nBùi Thị Diễm Hằng", size=12)
    set_cell_text(
        sign.rows[1].cells[1],
        "Học sinh 1: ................\n\nHọc sinh 2: ................",
        size=12,
    )

    add_para(doc, "PHÊ DUYỆT DỰ ÁN", 13, True, align="center", space_before=16)
    add_para(
        doc,
        "…………………………………………………………………………………………………………\n"
        "…………………………………………………………………………………………………………",
        12,
        align="left",
    )
    add_para(doc, "Hải Phòng, ngày ...... tháng ...... năm 2026", 13, align="right", space_before=8)
    add_para(doc, "HIỆU TRƯỞNG", 13, True, align="right")

    out = OUT / "Phu_luc_7_Ke_hoach_nghien_cuu_ThuyCanh_IoT.docx"
    doc.save(out)
    print("Saved", out)
    return out


if __name__ == "__main__":
    create_phu_luc_5()
    create_phu_luc_7()
    print("DONE")
    for f in sorted(OUT.glob("Phu_luc_*.docx")):
        print(f.name, f.stat().st_size)
